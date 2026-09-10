"""Tests for the resolved OSS inventory Excel export."""

from io import BytesIO
import unittest

from openpyxl import load_workbook

from inventory_export import inventory_xlsx


def sample_inventory() -> dict:
    return {
        "sources": ["maven-resolved"],
        "completeness": "complete",
        "excluded_modules": [],
        "entries": [
            {
                "groupId": "org.springframework",
                "artifactId": "spring-core",
                "version": "6.1.0",
                "purl": "pkg:maven/org.springframework/spring-core@6.1.0",
                "consumers": [{
                    "module": "g:app:1", "relationship": "direct",
                    "scopes": ["compile"], "paths": [["g:app:jar:1", "org.springframework:spring-core:jar:6.1.0"]],
                    "paths_truncated": False,
                }],
            },
            {
                "groupId": "org.springframework",
                "artifactId": "spring-beans",
                "version": "6.1.0",
                "purl": "pkg:maven/org.springframework/spring-beans@6.1.0",
                "consumers": [{
                    "module": "g:app:1", "relationship": "transitive",
                    "scopes": ["runtime"], "paths": [["g:app:jar:1", "x:mid:jar:1", "org.springframework:spring-beans:jar:6.1.0"]],
                    "paths_truncated": True,
                }],
            },
            {
                "groupId": "org.slf4j",
                "artifactId": "slf4j-api",
                "version": "2.0.0",
                "purl": "pkg:maven/org.slf4j/slf4j-api@2.0.0",
                "consumers": [{
                    "module": "g:app:1", "relationship": "direct",
                    "scopes": ["compile"], "paths": [["g:app:jar:1", "org.slf4j:slf4j-api:jar:2.0.0"]],
                    "paths_truncated": False,
                }],
            },
        ],
    }


class InventoryExportTests(unittest.TestCase):
    def test_export_has_summary_distribution_and_web_columns(self):
        workbook = load_workbook(BytesIO(inventory_xlsx(sample_inventory())))
        self.assertEqual(["Summary", "Inventory"], workbook.sheetnames)

        summary = workbook["Summary"]
        self.assertEqual("Resolved OSS Inventory Summary", summary["A1"].value)
        self.assertEqual("Total distinct external libraries", summary["A5"].value)
        self.assertEqual(3, summary["B5"].value)
        distribution = {
            summary.cell(row, 1).value: summary.cell(row, 2).value
            for row in range(11, summary.max_row + 1)
        }
        self.assertEqual(2, distribution["org.springframework"])
        self.assertEqual(1, distribution["org.slf4j"])

        detail = workbook["Inventory"]
        self.assertEqual(
            ["Coordinate", "Version", "Consumer", "Relationship", "Scopes", "Dependency path", "PURL"],
            [detail.cell(1, column).value for column in range(1, 8)],
        )
        self.assertEqual(3, detail.max_row - 1)
        self.assertIn("additional paths truncated", detail.cell(3, 6).value)


if __name__ == "__main__":
    unittest.main()
