"""Excel export for the resolved OSS inventory."""

from __future__ import annotations

from collections import Counter
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(sheet, row: int, end_column: int) -> None:
    for column in range(1, end_column + 1):
        cell = sheet.cell(row, column)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def inventory_xlsx(inventory: dict) -> bytes:
    """Return a readable workbook built from resolved external inventory data.

    ``entries`` are distinct canonical artifacts. The Inventory sheet keeps
    the same consumer-level columns as the web table; multiple bounded paths
    for one consumer remain together in its Dependency path cell.
    """
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    detail = workbook.create_sheet("Inventory")

    entries = inventory["entries"]
    group_counts = Counter(entry["groupId"] for entry in entries)
    relationship_count = sum(len(entry["consumers"]) for entry in entries)

    summary["A1"] = "Resolved OSS Inventory Summary"
    summary["A1"].font = Font(size=14, bold=True)
    summary_rows = [
        ("Scan source", ", ".join(inventory["sources"])),
        ("Completeness", inventory["completeness"]),
        ("Total distinct external libraries", len(entries)),
        ("Consumer relationships", relationship_count),
        ("Excluded modules", len(inventory["excluded_modules"])),
    ]
    for row, (label, value) in enumerate(summary_rows, 3):
        summary.cell(row, 1, label)
        summary.cell(row, 2, value)
        summary.cell(row, 1).font = Font(bold=True)

    distribution_header_row = 10
    summary.cell(distribution_header_row, 1, "Maven groupId")
    summary.cell(distribution_header_row, 2, "Distinct libraries")
    _style_header(summary, distribution_header_row, 2)
    for row, (group_id, count) in enumerate(
        sorted(group_counts.items(), key=lambda item: (-item[1], item[0])),
        distribution_header_row + 1,
    ):
        summary.cell(row, 1, group_id)
        summary.cell(row, 2, count)
    if group_counts:
        distribution_end = summary.max_row
        distribution = Table(
            displayName="GroupDistribution",
            ref=f"A{distribution_header_row}:B{distribution_end}",
        )
        distribution.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        summary.add_table(distribution)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 24
    summary.freeze_panes = "A3"

    headers = [
        "Coordinate",
        "Version",
        "Consumer",
        "Relationship",
        "Scopes",
        "Dependency path",
        "PURL",
    ]
    detail.append(headers)
    _style_header(detail, 1, len(headers))
    for entry in entries:
        coordinate = f"{entry['groupId']}:{entry['artifactId']}"
        for consumer in entry["consumers"]:
            paths = [" → ".join(path) for path in consumer["paths"]]
            if consumer["paths_truncated"]:
                paths.append("… (additional paths truncated)")
            detail.append([
                coordinate,
                entry["version"],
                consumer["module"],
                consumer["relationship"],
                ", ".join(consumer["scopes"]),
                "\n".join(paths),
                entry["purl"],
            ])
    if detail.max_row > 1:
        table = Table(displayName="ResolvedOssInventory",
                      ref=f"A1:G{detail.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        detail.add_table(table)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:G{max(detail.max_row, 1)}"
    widths = (35, 15, 30, 16, 18, 90, 50)
    for index, width in enumerate(widths, 1):
        detail.column_dimensions[chr(64 + index)].width = width
    for row in range(2, detail.max_row + 1):
        detail.cell(row, 6).alignment = Alignment(wrap_text=True,
                                                    vertical="top")

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
